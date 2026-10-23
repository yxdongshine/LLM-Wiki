#!/usr/bin/env python3
"""将 cmp_user、cmp_statistic_login_daily、aries_profiles_new 三张表填入入湖申请表。"""
import os
from datetime import datetime
from copy import copy
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

TEMPLATE = os.path.join(os.path.dirname(__file__), '数据入湖申请表模版.xlsx')
desktop = os.path.expanduser('~/Desktop')
today = datetime.now().strftime('%Y.%m.%d')
today_compact = datetime.now().strftime('%Y%m%d')

# ── 统一样式 ──
FONT_DATA = Font(name='Microsoft YaHei', size=11, bold=False, color='FF333333')
FONT_DATA_BOLD = Font(name='Microsoft YaHei', size=11, bold=True, color='FF333333')
ALIGN_CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
ALIGN_LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)
THIN_BORDER = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin'),
)

def style_cell(cell, font=FONT_DATA, alignment=ALIGN_LEFT):
    """统一设置单元格样式。"""
    cell.font = font
    cell.alignment = alignment
    cell.border = THIN_BORDER

wb = load_workbook(TEMPLATE)
ws = wb['数据入湖申请表单']

# ── 基础信息 ──
base_cells = {
    'B4': 'One蛇口', 'D4': '刘琨', 'F4': '周佳佳',
    'B5': today,
}
for ref, val in base_cells.items():
    cell = ws[ref]
    cell.value = val
    style_cell(cell, FONT_DATA, ALIGN_LEFT)

# ── 来源数据库信息 ──
db_cells = {
    'B9': 'MySQL', 'D9': '生产',
    'B10': '100.102.1.47', 'D10': '3308',
    'B11': 'cmp_member', 'D11': 'cmp_member',
    'B12': 'cmpmall',
    'B13': '周佳佳',
}
for ref, val in db_cells.items():
    cell = ws[ref]
    cell.value = val
    style_cell(cell, FONT_DATA, ALIGN_LEFT)

# ── 同步表数据 ──
# 表1: cmp_user
table1_data = {
    'B': 'cmp_user', 'C': '小程序用户信息表',
    'D': '小表', 'E': '小表',
    'F': """CREATE TABLE `cmp_member`.`cmp_user` (
  `id` varchar(32) CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci NOT NULL COMMENT '主键id',
  `open_id` varchar(32) CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci DEFAULT NULL COMMENT '微信openid',
  `union_id` varchar(32) CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci DEFAULT NULL COMMENT '微信unionId',
  `one_user_id` varchar(32) CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci NOT NULL COMMENT '用户中心用户id',
  `session_key` varchar(32) CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci DEFAULT NULL COMMENT '微信key，用于解密手机号信息',
  `reach_time` datetime DEFAULT NULL COMMENT '触达时间',
  `login_flag` tinyint NOT NULL DEFAULT '0' COMMENT '登录标识：0：已登录，1：退出登录',
  `unbind_flag` tinyint DEFAULT '0' COMMENT '解绑标识：0：正常绑定，1：解绑微信；2：解绑手机号',
  `last_login_time` timestamp NULL DEFAULT NULL COMMENT '最后登录时间',
  `wx_scene` int DEFAULT '0' COMMENT '微信场景值',
  `source_primary` int DEFAULT NULL COMMENT '业务注册来源——一级',
  `source_secondary` varchar(32) DEFAULT NULL COMMENT '业务注册来源——二级',
  `source_app` varchar(32) DEFAULT NULL COMMENT '来源系统',
  `source_biz_key` varchar(64) DEFAULT NULL COMMENT '来源系统:业务主键',
  `create_date` datetime NOT NULL DEFAULT CURRENT_TIMESTAMP COMMENT '创建时间',
  `created_by` varchar(64) NOT NULL DEFAULT 'system' COMMENT '创建人',
  `update_date` datetime DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP COMMENT '更新时间',
  `updated_by` varchar(64) DEFAULT '' COMMENT '更新人',
  `delete_flag` tinyint(1) NOT NULL DEFAULT '0' COMMENT '删除标志：1: 已删除 0：未删除',
  PRIMARY KEY (`id`),
  UNIQUE KEY `udx_one_user_id` (`one_user_id`,`delete_flag`),
  UNIQUE KEY `udx_open_id` (`open_id`,`delete_flag`),
  KEY `idx_union_id` (`union_id`)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci COMMENT='小程序用户信息表'""",
    'G': 'id', 'H': '否（逻辑删除 delete_flag）',
    'I': '增量', 'J': '每日一次', 'K': 'update_date',
    'L': '需要', 'M': '—', 'N': '实时更新',
}

# 表2: cmp_statistic_login_daily
table2_data = {
    'B': 'cmp_statistic_login_daily', 'C': '小程序每日登录统计表',
    'D': '小表', 'E': '小表',
    'F': """CREATE TABLE `cmp_member`.`cmp_statistic_login_daily` (
  `id` varchar(32) CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci NOT NULL COMMENT '主键ID',
  `open_id` varchar(64) CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci DEFAULT NULL COMMENT '微信openid',
  `one_user_id` varchar(32) CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci NOT NULL COMMENT '用户id',
  `login_type` int DEFAULT '0' COMMENT '登录类型：1：登入，2：登出',
  `login_date` datetime DEFAULT NULL COMMENT '修改时间',
  `login_num` int DEFAULT '0' COMMENT '每日登录次数',
  `first_login_time` datetime DEFAULT NULL COMMENT '每日第一次登录时间',
  `last_login_time` datetime DEFAULT NULL COMMENT '每日最后一次登录时间',
  `create_date` datetime NOT NULL DEFAULT CURRENT_TIMESTAMP COMMENT '创建时间',
  `created_by` varchar(64) CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci NOT NULL DEFAULT 'system' COMMENT '创建人',
  `update_date` datetime DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP COMMENT '更新时间',
  `updated_by` varchar(64) CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci DEFAULT '' COMMENT '更新人',
  `delete_flag` tinyint(1) NOT NULL DEFAULT '0' COMMENT '删除标志：1: 已删除 0：未删除',
  PRIMARY KEY (`id`),
  KEY `idx_one_user_id` (`one_user_id`),
  KEY `idx_login_date` (`login_date`)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci COMMENT='小程序每日登录统计表'""",
    'G': 'id', 'H': '否（逻辑删除 delete_flag）',
    'I': '增量', 'J': '每日一次', 'K': 'update_date',
    'L': '需要', 'M': '—', 'N': '定时更新(update_date)',
}

# 表3: aries_profiles_new
table3_data = {
    'B': 'aries_profiles_new', 'C': 'One蛇口小程序用户基础信息表',
    'D': '小表', 'E': '小表',
    'F': """CREATE TABLE `cmp_member`.`aries_profiles_new` (
  `id` varchar(36) CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci NOT NULL DEFAULT (uuid()) COMMENT '主键ID',
  `one_user_id` varchar(64) CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci NOT NULL COMMENT '用户id',
  `nickname` varchar(255) DEFAULT NULL COMMENT '昵称',
  `avatar` varchar(500) DEFAULT NULL COMMENT '头像',
  `mobile_prefix` varchar(32) DEFAULT NULL COMMENT '手机号码前缀',
  `mobile` varchar(32) DEFAULT NULL COMMENT '手机号',
  `cert_type` varchar(50) DEFAULT NULL COMMENT '证件类型',
  `cert_no` varchar(50) DEFAULT NULL COMMENT '证件号',
  `cert_name` varchar(50) DEFAULT NULL COMMENT '证件姓名',
  `gender` int DEFAULT NULL COMMENT '性别: 其他(0) 男(1) 女(2)',
  `birthday` datetime DEFAULT NULL COMMENT '出生日期',
  `register_time` datetime DEFAULT NULL COMMENT '注册时间',
  `status` int DEFAULT '0' COMMENT '状态: 冻结(1) 正常(0)',
  `org_id` varchar(32) DEFAULT NULL COMMENT '运营主体id',
  `org_name` varchar(255) DEFAULT NULL COMMENT '运营主体name',
  `membership_name` varchar(50) DEFAULT NULL COMMENT '所属招商荟名称',
  `membership_code` varchar(64) DEFAULT NULL COMMENT '所属招商荟Code',
  `register_city_code` varchar(32) DEFAULT NULL COMMENT '城市code',
  `register_city_name` varchar(32) DEFAULT NULL COMMENT '城市名称',
  `active_city_code` varchar(32) DEFAULT NULL COMMENT '活跃城市',
  `active_city_name` varchar(32) DEFAULT NULL COMMENT '城市名称',
  `province` varchar(255) DEFAULT NULL COMMENT '省ID',
  `city` varchar(255) DEFAULT NULL COMMENT '市ID',
  `detail_address` varchar(256) DEFAULT NULL COMMENT '详细地址',
  `role_type` varchar(30) DEFAULT NULL COMMENT '身份角色->1：业主，2：家属，3：租客，4：普通会员',
  `estate_projects` varchar(256) DEFAULT NULL COMMENT '地产项目名称，多个用逗号分割',
  `property_projects` varchar(256) DEFAULT NULL COMMENT '物业项目名称，多个用逗号分割',
  `owner_auth_status` tinyint unsigned NOT NULL DEFAULT '0' COMMENT '业主认证状态：0-未提交，1-待审核，2-已认证，3-已驳回',
  `owner_auth_date` datetime DEFAULT NULL COMMENT '业主认证通过时间',
  `level_id` varchar(64) DEFAULT NULL COMMENT '会员等级id',
  `level` varchar(64) DEFAULT NULL COMMENT '会员等级',
  `exist_member_card` tinyint DEFAULT NULL COMMENT '是否有门槛会员权益卡',
  `threshold_member_card_status` varchar(30) DEFAULT NULL COMMENT '门槛会员卡状态',
  `sync_flag` tinyint DEFAULT '0' COMMENT '同步老会员表状态：0:待同步，1：已同步',
  `create_date` datetime NOT NULL DEFAULT CURRENT_TIMESTAMP COMMENT '创建时间',
  `update_date` datetime NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP COMMENT '修改时间',
  `created_by` varchar(255) DEFAULT NULL COMMENT '创建人',
  `updated_by` varchar(255) DEFAULT NULL COMMENT '修改人',
  `delete_flag` int NOT NULL DEFAULT '0' COMMENT '逻辑删除标记',
  `extra_json` text COMMENT '扩展信息JSON格式',
  PRIMARY KEY (`id`),
  UNIQUE KEY `udx_one_user_id` (`one_user_id`),
  UNIQUE KEY `udx_mobile` (`mobile`),
  KEY `idx_register_time` (`register_time`),
  KEY `idx_org_id` (`org_id`),
  KEY `idx_register_city_code` (`register_city_code`),
  KEY `idx_update_date` (`update_date`)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci COMMENT='One蛇口小程序用户基础信息表'""",
    'G': 'id', 'H': '否（逻辑删除 delete_flag）',
    'I': '增量', 'J': '每日一次', 'K': 'update_date',
    'L': '需要', 'M': '—', 'N': '实时更新',
}

for row_num, data in [(18, table1_data), (19, table2_data), (20, table3_data)]:
    # 序号列（居中）
    cell_a = ws.cell(row=row_num, column=1)
    cell_a.value = row_num - 17
    style_cell(cell_a, FONT_DATA, ALIGN_CENTER)
    # 数据列
    for col_letter, val in data.items():
        col_idx = ord(col_letter) - ord('A') + 1
        cell = ws.cell(row=row_num, column=col_idx)
        cell.value = val
        style_cell(cell, FONT_DATA, ALIGN_LEFT)
    # 统一行高
    ws.row_dimensions[row_num].height = 29.1

# ── 其他信息 ──
for ref, val in [('B30', '尽快'), ('B31', '新增小程序用户、登录统计及用户基础信息表入湖，支撑用户行为分析')]:
    cell = ws[ref]
    cell.value = val
    style_cell(cell, FONT_DATA, ALIGN_LEFT)

# ── 授权信息（Row 36） ──
auth_data = {'A': 'One蛇口', 'B': 'cmsk_super_app', 'C': today, 'D': '每日一次', 'E': '无特别需求，入湖同学根据资源情况自行安排'}
for col_letter, val in auth_data.items():
    col_idx = ord(col_letter) - ord('A') + 1
    cell = ws.cell(row=36, column=col_idx)
    cell.value = val
    style_cell(cell, FONT_DATA, ALIGN_LEFT)

# ── 表结构说明模板 ──
ws3 = wb['表结构说明模板']
for mr in list(ws3.merged_cells.ranges):
    if mr.min_row >= 2:
        ws3.unmerge_cells(str(mr))
for r in range(2, 201):
    for c in range(1, 5):
        ws3.cell(row=r, column=c).value = None

# 表1字段: cmp_user
fields_cmp_user = [
    ('id', 'varchar(32)', '主键id'),
    ('open_id', 'varchar(32)', '微信openid'),
    ('union_id', 'varchar(32)', '微信unionId'),
    ('one_user_id', 'varchar(32)', '用户中心用户id'),
    ('session_key', 'varchar(32)', '微信key，用于解密手机号信息'),
    ('reach_time', 'datetime', '触达时间'),
    ('login_flag', 'tinyint', '登录标识：0：已登录，1：退出登录'),
    ('unbind_flag', 'tinyint', '解绑标识：0：正常绑定，1：解绑微信；2：解绑手机号'),
    ('last_login_time', 'timestamp', '最后登录时间'),
    ('wx_scene', 'int', '微信场景值'),
    ('source_primary', 'int', '业务注册来源——一级（0其他，1个人转发，2业态跳转，3自流量，4推广二维码）'),
    ('source_secondary', 'varchar(32)', '业务注册来源——二级'),
    ('source_app', 'varchar(32)', '来源系统（One蛇口小程序-UMOP，票务系统-TBS）'),
    ('source_biz_key', 'varchar(64)', '来源系统业务主键'),
    ('create_date', 'datetime', '创建时间'),
    ('created_by', 'varchar(64)', '创建人'),
    ('update_date', 'datetime', '更新时间'),
    ('updated_by', 'varchar(64)', '更新人'),
    ('delete_flag', 'tinyint(1)', '删除标志：1-已删除 0-未删除'),
]

# 表2字段: cmp_statistic_login_daily
fields_login_daily = [
    ('id', 'varchar(32)', '主键ID'),
    ('open_id', 'varchar(64)', '微信openid'),
    ('one_user_id', 'varchar(32)', '用户id'),
    ('login_type', 'int', '登录类型：1-登入，2-登出'),
    ('login_date', 'datetime', '修改时间'),
    ('login_num', 'int', '每日登录次数'),
    ('first_login_time', 'datetime', '每日第一次登录时间'),
    ('last_login_time', 'datetime', '每日最后一次登录时间'),
    ('create_date', 'datetime', '创建时间'),
    ('created_by', 'varchar(64)', '创建人'),
    ('update_date', 'datetime', '更新时间'),
    ('updated_by', 'varchar(64)', '更新人'),
    ('delete_flag', 'tinyint(1)', '删除标志：1-已删除 0-未删除'),
]

# 表3字段: aries_profiles_new
fields_aries_profiles = [
    ('id', 'varchar(36)', '主键ID'),
    ('one_user_id', 'varchar(64)', '用户id'),
    ('nickname', 'varchar(255)', '昵称'),
    ('avatar', 'varchar(500)', '头像'),
    ('mobile_prefix', 'varchar(32)', '手机号码前缀'),
    ('mobile', 'varchar(32)', '手机号'),
    ('cert_type', 'varchar(50)', '证件类型'),
    ('cert_no', 'varchar(50)', '证件号'),
    ('cert_name', 'varchar(50)', '证件姓名'),
    ('gender', 'int', '性别: 其他(0) 男(1) 女(2)'),
    ('birthday', 'datetime', '出生日期'),
    ('register_time', 'datetime', '注册时间'),
    ('status', 'int', '状态: 冻结(1) 正常(0)'),
    ('org_id', 'varchar(32)', '运营主体id'),
    ('org_name', 'varchar(255)', '运营主体name'),
    ('membership_name', 'varchar(50)', '所属招商荟名称'),
    ('membership_code', 'varchar(64)', '所属招商荟Code'),
    ('register_city_code', 'varchar(32)', '城市code'),
    ('register_city_name', 'varchar(32)', '城市名称'),
    ('active_city_code', 'varchar(32)', '活跃城市'),
    ('active_city_name', 'varchar(32)', '城市名称'),
    ('province', 'varchar(255)', '省ID'),
    ('city', 'varchar(255)', '市ID'),
    ('detail_address', 'varchar(256)', '详细地址'),
    ('role_type', 'varchar(30)', '身份角色(1业主 2家属 3租客 4普通会员)'),
    ('estate_projects', 'varchar(256)', '地产项目名称，多个用逗号分割'),
    ('property_projects', 'varchar(256)', '物业项目名称，多个用逗号分割'),
    ('owner_auth_status', 'tinyint unsigned', '业主认证状态(0未提交 1待审核 2已认证 3已驳回)'),
    ('owner_auth_date', 'datetime', '业主认证通过时间'),
    ('level_id', 'varchar(64)', '会员等级id'),
    ('level', 'varchar(64)', '会员等级'),
    ('exist_member_card', 'tinyint', '是否有门槛会员权益卡'),
    ('threshold_member_card_status', 'varchar(30)', '门槛会员卡状态'),
    ('sync_flag', 'tinyint', '同步老会员表状态(0待同步 1已同步)'),
    ('create_date', 'datetime', '创建时间'),
    ('update_date', 'datetime', '修改时间'),
    ('created_by', 'varchar(255)', '创建人'),
    ('updated_by', 'varchar(255)', '修改人'),
    ('delete_flag', 'int', '逻辑删除标记'),
    ('extra_json', 'text', '扩展信息JSON格式'),
]

def write_table_schema(ws, start_row, table_name, table_cn, fields):
    n = len(fields)
    if n > 1:
        ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row + n - 1, end_column=1)
    style_cell(ws.cell(row=start_row, column=1), FONT_DATA, ALIGN_LEFT)
    ws.cell(row=start_row, column=1).value = f'{table_name}\n（{table_cn}）'
    for i, (fn, ft, fd) in enumerate(fields):
        r = start_row + i
        style_cell(ws.cell(row=r, column=2), FONT_DATA, ALIGN_LEFT)
        style_cell(ws.cell(row=r, column=3), FONT_DATA, ALIGN_LEFT)
        style_cell(ws.cell(row=r, column=4), FONT_DATA, ALIGN_LEFT)
        ws.cell(row=r, column=2).value = fn
        ws.cell(row=r, column=3).value = ft
        ws.cell(row=r, column=4).value = fd
    return start_row + n

row = 2
row = write_table_schema(ws3, row, 'cmp_user', '小程序用户信息表', fields_cmp_user)
row = write_table_schema(ws3, row, 'cmp_statistic_login_daily', '小程序每日登录统计表', fields_login_daily)
row = write_table_schema(ws3, row, 'aries_profiles_new', 'One蛇口小程序用户基础信息表', fields_aries_profiles)

out = os.path.join(desktop, f'数据入湖申请表_One蛇口_{today_compact}.xlsx')
wb.save(out)
print(f'✅ 已生成: {out}')
