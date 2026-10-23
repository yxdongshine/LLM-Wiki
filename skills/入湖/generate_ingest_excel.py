#!/usr/bin/env python3
"""生成数据入湖申请表 Excel 文档。

基于模板 skills/入湖/数据入湖申请表模版.xlsx 复制并清空数据区，
保留所有样式和结构，并预填默认值。用法：python3 generate_ingest_excel.py [输出路径]
"""

import sys
import os
from datetime import datetime
from copy import copy
from openpyxl import load_workbook


TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), '数据入湖申请表模版.xlsx')

# ── 默认值（可在此修改） ──
DEFAULTS = {
    'project_group': 'One蛇口',       # 项目组
    'product_manager': '刘琨',         # 产品经理
    'tech_manager': '周佳佳',          # 技术经理
    'source_tech_manager': '周佳佳',   # 源库技术经理
}

# 数据入湖申请表单 sheet 中需要清空的单元格
FORM_CLEAR_CELLS = [
    # 基础申请信息
    (4, 2), (4, 4), (4, 6),  # 项目组, 产品经理, 技术经理
    (5, 2),                    # 申请日期
    # 来源数据库信息
    (9, 2), (9, 4),            # 数据库类型, 数据库环境
    (10, 2), (10, 4),          # 数据库地址, 端口
    (11, 2), (11, 4),          # 数据库名, Schema
    (12, 2),                   # 访问账号
    (13, 2),                   # 源库技术经理
    # 期望完成时间
    (30, 2),
    # 申请事由
    (31, 2),
    # 授权信息
    (36, 1), (36, 2), (36, 3), (36, 4), (36, 5),
]

# 表结构说明模板 sheet 中需要清空的单元格（保留表头）
SCHEMA_CLEAR_ROWS = range(2, 201)  # 清空第2行到200行


def clear_form_data(ws):
    """清空表单数据区，保留样式。"""
    for row, col in FORM_CLEAR_CELLS:
        cell = ws.cell(row=row, column=col)
        cell.value = None

    # 清空同步表数据区（Row 18-26, Col 1-14）
    for row_num in range(18, 27):
        for col_num in range(1, 15):
            cell = ws.cell(row=row_num, column=col_num)
            cell.value = None


def fill_defaults(ws):
    """预填默认值。"""
    today = datetime.now().strftime('%Y.%m.%d')

    defaults_map = {
        (4, 2): DEFAULTS['project_group'],      # 项目组
        (4, 4): DEFAULTS['product_manager'],     # 产品经理
        (4, 6): DEFAULTS['tech_manager'],        # 技术经理
        (5, 2): today,                            # 申请日期
        (13, 2): DEFAULTS['source_tech_manager'], # 源库技术经理
    }

    for (row, col), value in defaults_map.items():
        ws.cell(row=row, column=col).value = value


def clear_schema_data(ws):
    """清空表结构说明数据区，保留表头。"""
    # 先取消所有合并单元格（保留表头 Row 1）
    merged_to_unmerge = []
    for merged_range in ws.merged_cells.ranges:
        if merged_range.min_row >= 2:
            merged_to_unmerge.append(str(merged_range))
    for mr in merged_to_unmerge:
        ws.unmerge_cells(mr)

    for row_num in SCHEMA_CLEAR_ROWS:
        for col_num in range(1, 5):
            cell = ws.cell(row=row_num, column=col_num)
            cell.value = None


def generate(output_path=None):
    """生成数据入湖申请表 Excel（含默认值）。"""
    if not os.path.exists(TEMPLATE_PATH):
        print(f'❌ 模板文件不存在: {TEMPLATE_PATH}')
        return None

    wb = load_workbook(TEMPLATE_PATH)

    # 清空并填充数据入湖申请表单
    if '数据入湖申请表单' in wb.sheetnames:
        ws = wb['数据入湖申请表单']
        clear_form_data(ws)
        fill_defaults(ws)

    # 清空表结构说明模板
    if '表结构说明模板' in wb.sheetnames:
        clear_schema_data(wb['表结构说明模板'])

    # 保存（默认输出到桌面）
    if output_path is None:
        date_str = datetime.now().strftime('%Y%m%d')
        desktop = os.path.expanduser('~/Desktop')
        output_path = os.path.join(desktop, f'数据入湖申请表_{date_str}.xlsx')

    wb.save(output_path)
    print(f'✅ 数据入湖申请表已生成: {output_path}')
    return output_path


if __name__ == '__main__':
    path = sys.argv[1] if len(sys.argv) > 1 else None
    generate(path)