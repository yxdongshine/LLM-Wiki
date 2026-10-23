#!/usr/bin/env python3
"""将 activity_signin_result（签到结果表，新会员营销切换·变更4）填入入湖申请表。

默认值全部沿用 fill_cmp_member_tables.py / fill_business_entity.py 两次历史申请：
cmp_member 库 / 100.102.1.47:3308 / 账号 cmpmall / 增量+每日一次+update_date。
"""
import os
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, Border, Side

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, '数据入湖申请表模版.xlsx')
desktop = os.path.expanduser('~/Desktop')
today = datetime.now().strftime('%Y.%m.%d')
today_compact = datetime.now().strftime('%Y%m%d')

FONT_DATA = Font(name='Microsoft YaHei', size=11, bold=False, color='FF333333')
ALIGN_CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
ALIGN_LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)
THIN_BORDER = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin'),
)


def style_cell(cell, font=FONT_DATA, alignment=ALIGN_LEFT):
    cell.font = font
    cell.alignment = alignment
    cell.border = THIN_BORDER


wb = load_workbook(TEMPLATE)
ws = wb['数据入湖申请表单']

# ── 一、基础申请信息 ──
for ref, val in {'B4': 'One蛇口', 'D4': '刘琨', 'F4': '周佳佳', 'B5': today}.items():
    ws[ref].value = val
    style_cell(ws[ref])

# ── 三、来源数据库信息 ──
for ref, val in {
    'B9': 'MySQL', 'D9': '生产',
    'B10': '100.102.1.47', 'D10': '3308',
    'B11': 'cmp_member', 'D11': 'cmp_member',
    'B12': 'cmpmall',
    'B13': '周佳佳',
}.items():
    ws[ref].value = val
    style_cell(ws[ref])

# ── 四、同步表基本信息（Row 18，单表） ──
DDL = """CREATE TABLE `cmp_member`.`activity_signin_result` (
  `id` varchar(64) CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci NOT NULL COMMENT '主键ID',
  `activity_id` varchar(64) CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci NOT NULL COMMENT '活动ID',
  `member_id` varchar(64) CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci NOT NULL COMMENT '会员ID',
  `user_id` varchar(64) CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci DEFAULT NULL COMMENT 'ONE用户ID',
  `phone` varchar(20) CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci DEFAULT NULL COMMENT '手机号',
  `nick_name` varchar(255) CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci DEFAULT NULL COMMENT '会员昵称',
  `business_entity_id` varchar(64) CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci DEFAULT NULL COMMENT '运营主体id',
  `business_entity_name` varchar(255) CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci DEFAULT NULL COMMENT '运营主体name',
  `register_city_code` varchar(64) CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci DEFAULT NULL COMMENT '注册城市Code',
  `register_city_name` varchar(255) CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci DEFAULT NULL COMMENT '注册城市名称',
  `consecutive_sign_count` int DEFAULT '0' COMMENT '连续签到次数',
  `max_consecutive_count` int DEFAULT '0' COMMENT '最大连续签到次数',
  `total_sign_days` int DEFAULT '0' COMMENT '累计签到天数',
  `total_sign_count` int DEFAULT '0' COMMENT '累计签到次数',
  `latest_sign_time` datetime DEFAULT NULL COMMENT '最新签到时间',
  `cycle_start_time` datetime DEFAULT NULL COMMENT '周期开始时间',
  `cycle_end_time` datetime(3) DEFAULT NULL COMMENT '周期结束时间',
  `create_date` datetime NOT NULL DEFAULT CURRENT_TIMESTAMP COMMENT '创建时间',
  `update_date` datetime NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP COMMENT '更新时间',
  `created_by` varchar(64) CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci DEFAULT NULL COMMENT '创建人',
  `updated_by` varchar(64) CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci DEFAULT NULL COMMENT '更新人',
  `delete_flag` tinyint(1) NOT NULL DEFAULT '0' COMMENT '删除标识',
  PRIMARY KEY (`id`),
  UNIQUE KEY `uk_activity_member` (`activity_id`,`member_id`,`cycle_start_time`,`cycle_end_time`),
  KEY `idx_user_id` (`user_id`)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci COMMENT='签到结果表（用户签到统计）'"""

table1 = {
    'B': 'activity_signin_result',
    'C': '签到结果表（用户签到统计）',
    'D': '小表',
    'E': '小表',
    'F': DDL,
    'G': 'id（另有唯一键 uk_activity_member：activity_id+member_id+cycle_start_time+cycle_end_time）',
    'H': '否（逻辑删除 delete_flag）',
    'I': '增量',
    'J': '每日一次',
    'K': 'update_date',
    'L': '需要',
    'M': '—',
    'N': '定时更新(update_date)',
}

cell_a = ws.cell(row=18, column=1)
cell_a.value = 1
style_cell(cell_a, FONT_DATA, ALIGN_CENTER)
for col_letter, val in table1.items():
    cell = ws.cell(row=18, column=ord(col_letter) - ord('A') + 1)
    cell.value = val
    style_cell(cell)
ws.row_dimensions[18].height = 29.1

# ── 五、其他信息 ──
for ref, val in [
    ('B30', '尽快'),
    ('B31', '新会员营销切换（变更4·签到活动表替换）：新增 CMP_MEMBER 库签到结果表 activity_signin_result 入湖，'
            '替代原票务侧签到结果表 HOEST0154_CAMPAIGN_SIGN_RESULT_SS 作为签到数据源，支撑签到统计分析'),
]:
    ws[ref].value = val
    style_cell(ws[ref])

# ── 六、授权需求和抽数频率及时间 ──
auth = {'A': 'One蛇口', 'B': 'cmsk_super_app', 'C': today, 'D': '每日一次',
        'E': '无特别需求，入湖同学根据资源情况自行安排'}
for col_letter, val in auth.items():
    cell = ws.cell(row=36, column=ord(col_letter) - ord('A') + 1)
    cell.value = val
    style_cell(cell)

# ── 表结构说明模板 ──
ws3 = wb['表结构说明模板']
for mr in list(ws3.merged_cells.ranges):
    if mr.min_row >= 2:
        ws3.unmerge_cells(str(mr))
for r in range(2, 201):
    for c in range(1, 5):
        ws3.cell(row=r, column=c).value = None

fields = [
    ('id', 'varchar(64)', '主键ID'),
    ('activity_id', 'varchar(64)', '活动ID'),
    ('member_id', 'varchar(64)', '会员ID'),
    ('user_id', 'varchar(64)', 'ONE用户ID'),
    ('phone', 'varchar(20)', '手机号'),
    ('nick_name', 'varchar(255)', '会员昵称'),
    ('business_entity_id', 'varchar(64)', '运营主体id'),
    ('business_entity_name', 'varchar(255)', '运营主体name'),
    ('register_city_code', 'varchar(64)', '注册城市Code'),
    ('register_city_name', 'varchar(255)', '注册城市名称'),
    ('consecutive_sign_count', 'int', '连续签到次数'),
    ('max_consecutive_count', 'int', '最大连续签到次数'),
    ('total_sign_days', 'int', '累计签到天数'),
    ('total_sign_count', 'int', '累计签到次数'),
    ('latest_sign_time', 'datetime', '最新签到时间'),
    ('cycle_start_time', 'datetime', '周期开始时间'),
    ('cycle_end_time', 'datetime(3)', '周期结束时间'),
    ('create_date', 'datetime', '创建时间（DEFAULT CURRENT_TIMESTAMP）'),
    ('update_date', 'datetime', '更新时间（ON UPDATE CURRENT_TIMESTAMP，增量同步依据）'),
    ('created_by', 'varchar(64)', '创建人'),
    ('updated_by', 'varchar(64)', '更新人'),
    ('delete_flag', 'tinyint(1)', '删除标识（0-未删除 1-已删除）'),
]

n = len(fields)
ws3.merge_cells(start_row=2, start_column=1, end_row=2 + n - 1, end_column=1)
style_cell(ws3.cell(row=2, column=1), FONT_DATA, ALIGN_LEFT)
ws3.cell(row=2, column=1).value = 'activity_signin_result\n（签到结果表·用户签到统计）'
for i, (fn, ft, fd) in enumerate(fields):
    r = 2 + i
    for c in (2, 3, 4):
        style_cell(ws3.cell(row=r, column=c))
    ws3.cell(row=r, column=2).value = fn
    ws3.cell(row=r, column=3).value = ft
    ws3.cell(row=r, column=4).value = fd

out = os.path.join(desktop, f'数据入湖申请表_One蛇口_{today_compact}.xlsx')
wb.save(out)
print(f'已生成: {out}')
print(f'字段数: {n}')
