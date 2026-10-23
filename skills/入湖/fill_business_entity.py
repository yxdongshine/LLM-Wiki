#!/usr/bin/env python3
"""将 business_entity 和 business_entity_zsh_city_rel 两张表填入入湖申请表。"""
import os
from datetime import datetime
from copy import copy
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

TEMPLATE = os.path.join(os.path.dirname(__file__), '数据入湖申请表模版.xlsx')
desktop = os.path.expanduser('~/Desktop')
today = datetime.now().strftime('%Y.%m.%d')
today_compact = datetime.now().strftime('%Y%m%d')

# ── 统一样式 ──
FONT_DATA = Font(name='Microsoft YaHei', size=11, bold=False, color='FF333333')
FONT_DATA_BOLD = Font(name='Microsoft YaHei', size=11, bold=True, color='FF333333')
ALIGN_CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
ALIGN_LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)
THIN_BORDER = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin'),
)

def style_cell(cell, font=FONT_DATA, alignment=ALIGN_LEFT):
    """统一设置单元格样式。"""
    cell.font = font
    cell.alignment = alignment
    cell.border = THIN_BORDER

def style_row(ws, row_num, col_range, font=FONT_DATA, alignment=ALIGN_LEFT):
    """统一设置整行样式。"""
    for c in col_range:
        style_cell(ws.cell(row=row_num, column=c), font, alignment)

wb = load_workbook(TEMPLATE)
ws = wb['数据入湖申请表单']

# ── 基础信息（统一样式） ──
base_cells = {
    'B4': 'One蛇口', 'D4': '刘琨', 'F4': '周佳佳',
    'B5': today,
}
for ref, val in base_cells.items():
    cell = ws[ref]
    cell.value = val
    style_cell(cell, FONT_DATA, ALIGN_LEFT)

# ── 来源数据库信息 ──
db_cells = {
    'B9': 'MySQL', 'D9': '生产',
    'B10': '100.102.1.47', 'D10': '3308',
    'B11': 'cmp_member', 'D11': 'cmp_member',
    'B12': 'cmpmall',
    'B13': '周佳佳',
}
for ref, val in db_cells.items():
    cell = ws[ref]
    cell.value = val
    style_cell(cell, FONT_DATA, ALIGN_LEFT)

# ── 同步表数据（Row 18-19，统一所有列格式） ──
table1_data = {
    'B': 'business_entity', 'C': '运营主体',
    'D': '小表', 'E': '小表',
    'F': """CREATE TABLE `cmp_member`.`business_entity` (
  `id` varchar(64) NOT NULL DEFAULT 'uuid()' COMMENT '主键',
  `code` varchar(64) DEFAULT NULL COMMENT '编码',
  `name` varchar(255) DEFAULT NULL COMMENT '名称',
  `cmp_org_id` varchar(64) DEFAULT NULL COMMENT '关联cmp系统组织id',
  `cmp_org_name` varchar(255) DEFAULT NULL COMMENT '关联cmp系统组织name',
  `path_number` varchar(512) DEFAULT NULL COMMENT '组织路径编码',
  `path` varchar(1024) DEFAULT NULL COMMENT '组织路径名称',
  `business_id` varchar(64) DEFAULT NULL COMMENT '归属业态id',
  `business_name` varchar(255) DEFAULT NULL COMMENT '归属业态名称',
  `create_date` datetime DEFAULT NULL COMMENT '创建时间',
  `update_date` datetime DEFAULT NULL COMMENT '修改时间',
  `created_by` varchar(64) DEFAULT NULL COMMENT '创建人',
  `updated_by` varchar(64) DEFAULT NULL COMMENT '修改人',
  `delete_flag` tinyint DEFAULT NULL COMMENT '逻辑删除标记(0-正常 1-已删除)',
  PRIMARY KEY (`id`)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COMMENT='运营主体'""",
    'G': 'id', 'H': '否（逻辑删除 delete_flag）',
    'I': '增量', 'J': '每日一次', 'K': 'update_date',
    'L': '需要', 'M': '—', 'N': '定时更新(update_date)',
}
table2_data = {
    'B': 'business_entity_zsh_city_rel', 'C': '运营主体-招商荟-城市关联关系',
    'D': '小表', 'E': '小表',
    'F': """CREATE TABLE `cmp_member`.`business_entity_zsh_city_rel` (
  `id` varchar(64) NOT NULL DEFAULT 'uuid()' COMMENT '主键',
  `business_entity_id` varchar(64) DEFAULT NULL COMMENT '运营主体ID',
  `business_entity_code` varchar(64) DEFAULT NULL COMMENT '运营主体编码',
  `zsh_id` varchar(64) DEFAULT NULL COMMENT '招商荟ID',
  `zsh_code` varchar(64) DEFAULT NULL COMMENT '招商荟编码',
  `city_id` varchar(64) DEFAULT NULL COMMENT '城市ID',
  `city_code` varchar(64) DEFAULT NULL COMMENT '城市编码',
  `create_date` datetime DEFAULT NULL COMMENT '创建时间',
  `update_date` datetime DEFAULT NULL COMMENT '修改时间',
  `created_by` varchar(64) DEFAULT NULL COMMENT '创建人',
  `updated_by` varchar(64) DEFAULT NULL COMMENT '修改人',
  `delete_flag` tinyint DEFAULT NULL COMMENT '逻辑删除标记(0-正常 1-已删除)',
  PRIMARY KEY (`id`)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COMMENT='运营主体-招商荟-城市关联关系'""",
    'G': 'id', 'H': '否（逻辑删除 delete_flag）',
    'I': '增量', 'J': '每日一次', 'K': 'update_date',
    'L': '需要', 'M': '—', 'N': '定时更新(update_date)',
}

for row_num, data in [(18, table1_data), (19, table2_data)]:
    # 序号列（居中）
    cell_a = ws.cell(row=row_num, column=1)
    cell_a.value = row_num - 17
    style_cell(cell_a, FONT_DATA, ALIGN_CENTER)
    # 数据列
    for col_letter, val in data.items():
        col_idx = ord(col_letter) - ord('A') + 1
        cell = ws.cell(row=row_num, column=col_idx)
        cell.value = val
        style_cell(cell, FONT_DATA, ALIGN_LEFT)
    # 统一行高
    ws.row_dimensions[row_num].height = 29.1

# ── 其他信息 ──
for ref, val in [('B30', '尽快'), ('B31', '新增运营主体及关联关系表入湖，支撑招商荟城市维度数据分析')]:
    cell = ws[ref]
    cell.value = val
    style_cell(cell, FONT_DATA, ALIGN_LEFT)

# ── 授权信息（Row 36，统一格式） ──
auth_data = {'A': 'One蛇口', 'B': 'cmsk_super_app', 'C': today, 'D': '每日一次', 'E': '无特别需求，入湖同学根据资源情况自行安排'}
for col_letter, val in auth_data.items():
    col_idx = ord(col_letter) - ord('A') + 1
    cell = ws.cell(row=36, column=col_idx)
    cell.value = val
    style_cell(cell, FONT_DATA, ALIGN_LEFT)

# ── 表结构说明模板 ──
ws3 = wb['表结构说明模板']
for mr in list(ws3.merged_cells.ranges):
    if mr.min_row >= 2:
        ws3.unmerge_cells(str(mr))
for r in range(2, 201):
    for c in range(1, 5):
        ws3.cell(row=r, column=c).value = None

fields1 = [
    ('id', 'varchar(64)', '主键'), ('code', 'varchar(64)', '编码'),
    ('name', 'varchar(255)', '名称'), ('cmp_org_id', 'varchar(64)', '关联cmp系统组织id'),
    ('cmp_org_name', 'varchar(255)', '关联cmp系统组织name'),
    ('path_number', 'varchar(512)', '组织路径编码（组织id以/分隔，不含根0）'),
    ('path', 'varchar(1024)', '组织路径名称（组织名称以/分隔）'),
    ('business_id', 'varchar(64)', '归属业态id'), ('business_name', 'varchar(255)', '归属业态名称'),
    ('create_date', 'datetime', '创建时间'), ('update_date', 'datetime', '修改时间'),
    ('created_by', 'varchar(64)', '创建人'), ('updated_by', 'varchar(64)', '修改人'),
    ('delete_flag', 'tinyint', '逻辑删除标记(0-正常 1-已删除)'),
]
fields2 = [
    ('id', 'varchar(64)', '主键'), ('business_entity_id', 'varchar(64)', '运营主体ID'),
    ('business_entity_code', 'varchar(64)', '运营主体编码'),
    ('zsh_id', 'varchar(64)', '招商荟ID'), ('zsh_code', 'varchar(64)', '招商荟编码'),
    ('city_id', 'varchar(64)', '城市ID'), ('city_code', 'varchar(64)', '城市编码'),
    ('create_date', 'datetime', '创建时间'), ('update_date', 'datetime', '修改时间'),
    ('created_by', 'varchar(64)', '创建人'), ('updated_by', 'varchar(64)', '修改人'),
    ('delete_flag', 'tinyint', '逻辑删除标记(0-正常 1-已删除)'),
]

def write_table_schema(ws, start_row, table_name, table_cn, fields):
    n = len(fields)
    if n > 1:
        ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row + n - 1, end_column=1)
    style_cell(ws.cell(row=start_row, column=1), FONT_DATA, ALIGN_LEFT)
    ws.cell(row=start_row, column=1).value = f'{table_name}\n（{table_cn}）'
    for i, (fn, ft, fd) in enumerate(fields):
        r = start_row + i
        style_cell(ws.cell(row=r, column=2), FONT_DATA, ALIGN_LEFT)
        style_cell(ws.cell(row=r, column=3), FONT_DATA, ALIGN_LEFT)
        style_cell(ws.cell(row=r, column=4), FONT_DATA, ALIGN_LEFT)
        ws.cell(row=r, column=2).value = fn
        ws.cell(row=r, column=3).value = ft
        ws.cell(row=r, column=4).value = fd
    return start_row + n

row = 2
row = write_table_schema(ws3, row, 'business_entity', '运营主体', fields1)
row = write_table_schema(ws3, row, 'business_entity_zsh_city_rel', '运营主体-招商荟-城市关联关系', fields2)

out = os.path.join(desktop, f'数据入湖申请表_{today_compact}.xlsx')
wb.save(out)
print(f'✅ 已生成: {out}')